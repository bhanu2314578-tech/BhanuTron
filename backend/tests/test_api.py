import unittest
from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4

from fastapi.testclient import TestClient
from docx import Document as WordDocument

from main import app
from routers.auth import AuthUser, get_current_user
from routers.documents import _detect_file_type
from services.document_processor import process_document


class FakeQuery:
    def __init__(self, database: "FakeSupabase", table: str) -> None:
        self.database = database
        self.table_name = table
        self.operation = "select"
        self.payload = None
        self.filters: list[tuple[str, str, object]] = []
        self.single = False
        self.limit_count: int | None = None

    def select(self, *_args: object, **_kwargs: object) -> "FakeQuery":
        del _args, _kwargs
        return self

    def insert(self, payload: object) -> "FakeQuery":
        self.operation = "insert"
        self.payload = payload
        return self

    def update(self, payload: dict[str, object]) -> "FakeQuery":
        self.operation = "update"
        self.payload = payload
        return self

    def delete(self) -> "FakeQuery":
        self.operation = "delete"
        return self

    def eq(self, key: str, value: object) -> "FakeQuery":
        self.filters.append(("eq", key, value))
        return self

    def in_(self, key: str, value: object) -> "FakeQuery":
        self.filters.append(("in", key, value))
        return self

    def order(self, *_args: object, **_kwargs: object) -> "FakeQuery":
        del _args, _kwargs
        return self

    def limit(self, value: int) -> "FakeQuery":
        self.limit_count = value
        return self

    def maybe_single(self) -> "FakeQuery":
        self.single = True
        return self

    def _matches(self, row: dict[str, object]) -> bool:
        for operator, key, expected in self.filters:
            if operator == "eq" and row.get(key) != expected:
                return False
            if operator == "in" and row.get(key) not in expected:
                return False
        return True

    def execute(self) -> SimpleNamespace:
        rows = self.database.rows.setdefault(self.table_name, [])
        if self.operation == "insert":
            values = self.payload if isinstance(self.payload, list) else [self.payload]
            inserted = []
            for value in values:
                row = dict(value)
                row.setdefault("id", str(uuid4()))
                now = datetime.now(timezone.utc).isoformat()
                row.setdefault("created_at", now)
                row.setdefault("updated_at", now)
                rows.append(row)
                inserted.append(row)
            return SimpleNamespace(data=inserted)

        matching = [row for row in rows if self._matches(row)]
        if self.operation == "update":
            for row in matching:
                row.update(self.payload)
            return SimpleNamespace(data=matching)
        if self.operation == "delete":
            self.database.rows[self.table_name] = [
                row for row in rows if not self._matches(row)
            ]
            return SimpleNamespace(data=matching)
        if self.limit_count is not None:
            matching = matching[: self.limit_count]
        return SimpleNamespace(data=(matching[0] if self.single and matching else None) if self.single else matching)


class FakeStorageBucket:
    def upload(self, *_args: object, **_kwargs: object) -> None:
        del _args, _kwargs
        return None

    def remove(self, *_args: object, **_kwargs: object) -> None:
        del _args, _kwargs
        return None


class FakeSupabase:
    def __init__(self) -> None:
        self.rows: dict[str, list[dict[str, object]]] = {}
        self.storage = SimpleNamespace(from_=lambda _: FakeStorageBucket())

    def table(self, table_name: str) -> FakeQuery:
        return FakeQuery(self, table_name)


class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.client = TestClient(app)

    def setUp(self) -> None:
        self.user = AuthUser(id=str(uuid4()), email="person@example.com", name="Person")
        app.dependency_overrides[get_current_user] = lambda: self.user

    def tearDown(self) -> None:
        app.dependency_overrides.pop(get_current_user, None)

    def test_health_check_does_not_require_credentials(self) -> None:
        response = self.client.get("/health")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"status": "ok"})

    def test_protected_route_requires_bearer_token(self) -> None:
        override = app.dependency_overrides.pop(get_current_user)
        try:
            response = self.client.get("/api/documents")
        finally:
            app.dependency_overrides[get_current_user] = override

        self.assertEqual(response.status_code, 401)
        self.assertEqual(response.json()["detail"], "Authentication required")

    def test_auth_route_reports_missing_backend_configuration(self) -> None:
        from fastapi import HTTPException

        with patch(
            "routers.auth.get_auth_client",
            side_effect=HTTPException(status_code=503, detail="Authentication service is not configured"),
        ):
            response = self.client.post(
                "/api/auth/login",
                json={"email": "person@example.com", "password": "not-a-real-password"},
            )

        self.assertEqual(response.status_code, 503)
        self.assertNotIn("not-a-real-password", response.text)

    def test_sign_in_returns_supabase_access_token(self) -> None:
        auth = SimpleNamespace(
            sign_in_with_password=lambda _: SimpleNamespace(
                session=SimpleNamespace(access_token="test-access-token"),
                user=SimpleNamespace(
                    id="user-id",
                    email="person@example.com",
                    user_metadata={"name": "Person"},
                ),
            )
        )
        with patch("routers.auth.get_auth_client", return_value=SimpleNamespace(auth=auth)):
            response = self.client.post(
                "/api/auth/login",
                json={"email": "person@example.com", "password": "example-password"},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["token"], "test-access-token")
        self.assertEqual(response.json()["user"]["name"], "Person")

    def test_chat_persists_messages_and_returns_sources(self) -> None:
        database = FakeSupabase()
        sources = [
            {
                "id": "chunk-id",
                "document": "guide.txt",
                "page": 1,
                "section": "Page 1",
                "snippet": "A relevant excerpt",
            }
        ]
        with (
            patch("routers.chat.get_supabase", return_value=database),
            patch(
                "routers.chat.answer_question",
                return_value=("Answer from document", sources, 12),
            ),
        ):
            response = self.client.post(
                "/api/chat",
                json={"question": "What does it say?"},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"answer": "Answer from document", "sources": sources})
        self.assertEqual(len(database.rows["messages"]), 2)
        self.assertEqual(database.rows["messages"][0]["role"], "user")
        self.assertEqual(database.rows["messages"][1]["role"], "assistant")

    def test_upload_processes_plain_text_and_returns_document(self) -> None:
        database = FakeSupabase()
        with (
            patch("routers.documents.get_supabase", return_value=database),
            patch("routers.documents.insert_chunks"),
        ):
            response = self.client.post(
                "/api/upload",
                files={"file": ("guide.txt", b"Document text for retrieval.", "text/plain")},
            )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["name"], "guide.txt")
        self.assertEqual(response.json()["status"], "processed")
        self.assertEqual(database.rows["documents"][0]["user_id"], self.user.id)

    def test_upload_alias_is_registered(self) -> None:
        paths = app.openapi()["paths"]

        self.assertIn("/api/upload", paths)
        self.assertIn("/api/documents", paths)

    def test_swagger_ui_and_openapi_are_available(self) -> None:
        docs_response = self.client.get("/docs")
        schema_response = self.client.get("/openapi.json")

        self.assertEqual(docs_response.status_code, 200)
        self.assertIn("swagger-ui", docs_response.text)
        self.assertEqual(schema_response.status_code, 200)
        self.assertEqual(schema_response.json()["openapi"], "3.1.0")

    def test_every_api_operation_is_described_and_protected_routes_use_bearer_auth(self) -> None:
        schema = self.client.get("/openapi.json").json()
        operations = [
            (path, method, operation)
            for path, path_item in schema["paths"].items()
            for method, operation in path_item.items()
            if method in {"get", "post", "patch", "delete"}
        ]

        self.assertGreaterEqual(len(operations), 18)
        for path, method, operation in operations:
            with self.subTest(path=path, method=method):
                self.assertTrue(operation.get("summary"))
                self.assertTrue(operation.get("description"))
        self.assertEqual(
            schema["paths"]["/api/documents"]["get"]["security"],
            [{"HTTPBearer": []}],
        )
        self.assertIn("HTTPBearer", schema["components"]["securitySchemes"])

    def test_document_type_comes_from_filename_extension(self) -> None:
        class Upload:
            filename = r"..\private\report.PDF"

        self.assertEqual(_detect_file_type(Upload()), "pdf")

    def test_text_document_is_chunked(self) -> None:
        chunks, page_count = process_document(
            b"First line of the document.\nSecond line.",
            "txt",
        )

        self.assertEqual(page_count, 1)
        self.assertEqual(len(chunks), 1)
        self.assertIn("First line", chunks[0].content)

    def test_empty_text_document_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "No readable text"):
            process_document(b"  \n  ", "txt")

    def test_docx_table_text_is_retrievable(self) -> None:
        document = WordDocument()
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Table content for retrieval"
        buffer = BytesIO()
        document.save(buffer)

        chunks, page_count = process_document(buffer.getvalue(), "docx")

        self.assertEqual(page_count, 1)
        self.assertIn("Table content", chunks[0].content)


if __name__ == "__main__":
    unittest.main()
