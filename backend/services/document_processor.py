from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


@dataclass(frozen=True)
class DocumentChunk:
    content: str
    page_number: int
    section: str
    chunk_index: int


def _split_text(text: str, page_number: int, start_index: int) -> list[DocumentChunk]:
    normalized = " ".join(text.split())
    if not normalized:
        return []

    chunk_size = 1_000
    overlap = 150
    chunks: list[DocumentChunk] = []
    start = 0
    while start < len(normalized):
        end = min(start + chunk_size, len(normalized))
        if end < len(normalized):
            boundary = normalized.rfind(" ", start, end)
            if boundary > start + chunk_size // 2:
                end = boundary
        content = normalized[start:end].strip()
        if content:
            chunks.append(
                DocumentChunk(
                    content=content,
                    page_number=page_number,
                    section=f"Page {page_number}",
                    chunk_index=start_index + len(chunks),
                )
            )
        if end == len(normalized):
            break
        start = max(end - overlap, start + 1)
    return chunks


def process_document(file_bytes: bytes, file_type: str) -> tuple[list[DocumentChunk], int]:
    if file_type == "pdf":
        try:
            reader = PdfReader(BytesIO(file_bytes))
            pages = [(index + 1, page.extract_text() or "") for index, page in enumerate(reader.pages)]
        except (PdfReadError, ValueError, OSError) as exc:
            raise ValueError("The PDF could not be read") from exc
        page_count = len(pages)
    elif file_type == "docx":
        try:
            document = Document(BytesIO(file_bytes))
        except (
            BadZipFile,
            KeyError,
            PackageNotFoundError,
            ValueError,
            OSError,
            XMLSyntaxError,
        ) as exc:
            raise ValueError("The DOCX file could not be read") from exc
        text_parts = [paragraph.text for paragraph in document.paragraphs]
        text_parts.extend(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        text = "\n".join(text_parts)
        pages = [(1, text)]
        page_count = 1 if text.strip() else 0
    elif file_type == "txt":
        try:
            text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("The text file must use UTF-8 encoding") from exc
        pages = [(1, text)]
        page_count = 1 if text.strip() else 0
    else:
        raise ValueError("Unsupported document type")

    chunks: list[DocumentChunk] = []
    for page_number, page_text in pages:
        chunks.extend(_split_text(page_text, page_number, len(chunks)))
    if not chunks:
        raise ValueError("No readable text was found in the document")
    return chunks, page_count
